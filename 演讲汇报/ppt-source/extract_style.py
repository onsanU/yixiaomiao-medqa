# -*- coding: utf-8 -*-
"""扒老师课程PPT的风格: 幻灯片数 / 背景色 / 字体色频次 / 字体名 / 每页标题"""
from pptx import Presentation
from pptx.util import Emu
from collections import Counter

PATH = "/mnt/d/Project/泰迪/deepseek医疗专家/基于DeepSeek大模型微调的医疗专家模型.pptx"
prs = Presentation(PATH)
print(f"幻灯片尺寸: {prs.slide_width/914400:.2f} x {prs.slide_height/914400:.2f} 英寸")
print(f"总页数: {len(prs.slides)}\n")

bg_counter = Counter()
color_counter = Counter()
font_counter = Counter()
size_counter = Counter()

for i, slide in enumerate(prs.slides, 1):
    # 背景色
    try:
        fill = slide.background.fill
        if fill.type is not None and str(fill.type) != "None":
            bg_counter[str(fill.fore_color.rgb)] += 1
    except Exception:
        pass
    texts = []
    for sh in slide.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text.strip().replace("\n", " / ")
            if t:
                texts.append(t[:60])
            for p in sh.text_frame.paragraphs:
                for r in p.runs:
                    if r.font.color and r.font.color.type is not None:
                        try:
                            color_counter[str(r.font.color.rgb)] += 1
                        except Exception:
                            pass
                    if r.font.name:
                        font_counter[r.font.name] += 1
                    if r.font.size:
                        size_counter[int(r.font.size.pt)] += 1
    print(f"[{i:02d}] {' | '.join(texts[:3])}")

print("\n=== 背景色频次 ===")
for c, n in bg_counter.most_common():
    print(f"  #{c}  x{n}")
print("\n=== 字体色 TOP12 ===")
for c, n in color_counter.most_common(12):
    print(f"  #{c}  x{n}")
print("\n=== 字体名 TOP6 ===")
for c, n in font_counter.most_common(6):
    print(f"  {c}  x{n}")
print("\n=== 字号 TOP10 ===")
for c, n in size_counter.most_common(10):
    print(f"  {c}pt  x{n}")
